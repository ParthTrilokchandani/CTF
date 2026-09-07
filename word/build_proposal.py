"""
Builds CTF_Client_Proposal.docx - a client-facing proposal document.

Contains no scripts, commands, credentials, real flags, exact exploitation
steps, or implementation details - management/client level only.

Usage: python3 build_proposal.py
Requires: python-docx, Pillow (for the diagrams, generated separately by
diagrams/gen_proposal_diagrams.py - run that first).
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS = os.path.join(OUT_DIR, "diagrams")

ACCENT = RGBColor(0x7A, 0x1F, 0x2B)
ACCENT2 = RGBColor(0x1E, 0x3A, 0x5F)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LIGHT_BG = "F4F2EC"


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def shade_paragraph(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_hr(doc, color="B8AFA0"):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)
    p.paragraph_format.space_after = Pt(14)
    return p


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def add_toc(doc):
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
    r = doc.add_paragraph()
    rr = r.add_run("(Right-click this area in Word and choose “Update Field” to generate the table of contents.)")
    rr.italic = True
    rr.font.size = Pt(9.5)
    rr.font.color.rgb = MUTED


def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Calibri Light"
    title.font.size = Pt(34)
    title.font.color.rgb = ACCENT

    for lvl, size in ((1, 17), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT if lvl == 1 else ACCENT2
        st.font.bold = True
        st.paragraph_format.space_before = Pt(20 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True


def add_page_number(paragraph, prefix=""):
    if prefix:
        r = paragraph.add_run(prefix)
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
    add_field(paragraph, "PAGE")
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def setup_header_footer(section, title_text):
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(title_text)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    r.bold = True
    r2 = hp.add_run("   |   CONFIDENTIAL - CLIENT PROPOSAL")
    r2.font.size = Pt(9)
    r2.font.color.rgb = MUTED
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "B8AFA0")
    borders.append(bottom)
    pPr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp, prefix="Page ")


def add_table(doc, headers, rows, col_widths=None, style="Light Grid Accent 2"):
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = style
    except KeyError:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_figure(doc, img_filename, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(os.path.join(DIAGRAMS, img_filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    cap.paragraph_format.space_after = Pt(14)


def note_box(doc, text):
    p = doc.add_paragraph()
    shade_paragraph(p, "FBF3E0")
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), "D8B96A")
        borders.append(el)
    pPr.append(borders)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x6E, 0x4A, 0x08)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, size=None, color=None, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    r.bold = bold
    r.italic = italic
    if align:
        p.alignment = align
    return p


# ---------------------------------------------------------------------------
def build():
    doc = Document()
    set_base_styles(doc)

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    # ---------------- Cover page ----------------
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COMPLIANCE OFFENSIVE ENGAGEMENT")
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    r.bold = True

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.add_run("Cybersecurity Capture-The-Flag\nEngagement Proposal")
    tr.font.size = Pt(32)
    tr.font.name = "Calibri Light"
    tr.font.color.rgb = ACCENT
    tr.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("A realistic, progressive, five-flag offensive-security exercise\nin a fully isolated, controlled virtual environment")
    sr.font.size = Pt(13)
    sr.italic = True
    sr.font.color.rgb = MUTED

    for _ in range(6):
        doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ["Prepared for: [Client Organization]", "Prepared by: [Your Organization]", "Date: [Proposal Date]"]:
        r = meta.add_run(line + "\n")
        r.font.size = Pt(11)
        r.font.color.rgb = INK

    doc.add_paragraph()
    conf = doc.add_paragraph()
    conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = conf.add_run("CONFIDENTIAL - Prepared exclusively for the recipient client. Not for redistribution.")
    cr.font.size = Pt(9)
    cr.italic = True
    cr.font.color.rgb = MUTED

    doc.add_page_break()

    # ---------------- Table of contents ----------------
    heading(doc, "Table of Contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    setup_header_footer(section, "Cybersecurity CTF Engagement Proposal")

    # ---------------- Executive Summary ----------------
    heading(doc, "1. Executive Summary")
    para(doc,
         "This proposal outlines a purpose-built cybersecurity Capture-The-Flag (CTF) "
         "exercise designed to give participants realistic, hands-on offensive-security "
         "practice in a fully controlled and isolated environment. The exercise is built "
         "around a believable intelligence/security-themed storyline rather than an "
         "obvious series of disconnected puzzles, so participants practice the same "
         "investigative mindset and methodology they would apply in a real engagement.")
    para(doc,
         "The engagement centers on a single target system containing five progressively "
         "difficult flags. Participants start with nothing but a target IP address and must "
         "apply reconnaissance, web analysis, steganography, credential handling, Linux "
         "enumeration, and controlled privilege escalation to reach full compromise. The "
         "environment is self-contained, safe to run repeatedly, and has no dependency on "
         "or access to production infrastructure.")

    # ---------------- CTF Objectives ----------------
    heading(doc, "2. CTF Objectives")
    bullets(doc, [
        "Provide hands-on practice spanning the full offensive engagement lifecycle: "
        "reconnaissance, initial access, lateral discovery, privilege escalation, and full compromise.",
        "Reinforce structured methodology, patience, and verification habits rather than "
        "rewarding guesswork.",
        "Present a believable, realistic scenario that keeps participants in an investigative "
        "mindset instead of an obvious “game” framing.",
        "Exercise a broad and representative skill set: network scanning, web enumeration, "
        "steganography, wordlist-based credential recovery, SSH usage, Linux permissions "
        "analysis, and basic reverse engineering.",
        "Deliver the exercise in a safe, isolated, and fully resettable environment suitable "
        "for repeated use across multiple sessions, teams, or cohorts.",
    ])

    # ---------------- Proposed Environment ----------------
    heading(doc, "3. Proposed Environment")
    para(doc,
         "The exercise is hosted entirely within a controlled Proxmox virtual environment, "
         "keeping every component - participant machines and the CTF target alike - "
         "isolated from the client's production systems.")
    add_table(doc, ["Component", "Description"], [
        ["Hosting platform", "Proxmox VE (virtualization) - dedicated to the exercise"],
        ["Participant machine(s)", "Kali Linux / Parrot Security attacker VMs, example network 10.10.x.x"],
        ["CTF target", "Dedicated Ubuntu Server VM hosting the exercise, example network 10.20.x.x"],
        ["Network isolation", "Dedicated virtual network segment with no route to production infrastructure"],
    ], col_widths=[2.0, 4.3])
    note_box(doc,
             "Both 10.10.x.x and 10.20.x.x above are illustrative examples only. Both ranges "
             "are fully configurable and will be adapted to fit the client's actual environment "
             "and addressing scheme.")

    # ---------------- Architecture ----------------
    heading(doc, "4. Architecture")
    para(doc,
         "At a high level, participant machines and the CTF target both run as virtual "
         "machines under Proxmox, connected only through an isolated network segment "
         "dedicated to the exercise. No component of the exercise requires or is granted "
         "access to the client's production network, management interfaces, or other "
         "virtual machines.")
    add_figure(doc, "proposal_architecture.png", "Figure 1: Proposed environment architecture")

    # ---------------- CTF Storyline ----------------
    heading(doc, "5. CTF Storyline")
    para(doc,
         "The exercise is framed around a fictional intelligence/security consultancy whose "
         "public-facing website looks entirely ordinary at first glance. Nothing on the "
         "surface announces that this is a CTF - there is no obvious “start here” "
         "messaging, no flag placeholders, and no puzzle-like styling. Participants are "
         "expected to treat the target the way they would a real reconnaissance target: "
         "methodically, and without assuming anything is exactly as it appears.")
    para(doc,
         "As the engagement progresses, participants uncover a trail of operational history "
         "the organization never properly cleaned up: forgotten credentials, an old backup "
         "location, a second account with elevated access, and finally a narrowly-scoped, "
         "intentionally controlled path to full administrative control. The storyline exists "
         "to keep participants oriented and motivated - it does not change the technical "
         "skills being tested, but it makes practicing them feel like a real investigation "
         "rather than an artificial exercise.")

    # ---------------- Complete Attack/Challenge Flow ----------------
    heading(doc, "6. Complete Attack/Challenge Flow")
    para(doc,
         "The exercise follows a single, internally consistent path from initial "
         "reconnaissance through to full compromise. Each stage below is described at a "
         "conceptual level; no exact commands, file locations, or credentials are included "
         "in this document.")
    numbered(doc, [
        "Participants receive only the CTF target's IP address and perform standard "
        "reconnaissance to identify running services.",
        "Initial web reconnaissance uncovers the organization's website and, through "
        "standard enumeration, a restricted path referenced by the site's crawler-control "
        "file.",
        "Following that path leads to a page that is not linked from normal navigation, "
        "where the first flag is present in an encoded form after reviewing the page's "
        "content.",
        "Separately, an image published on the website is found to contain hidden data, "
        "recoverable through steganographic analysis once the correct passphrase is "
        "identified via a wordlist-based approach.",
        "The recovered hidden data contains an encoded credential which, once decoded, "
        "grants SSH access to a first, intentionally low-privileged account.",
        "Enumeration of that account's home directory reveals a second encoded flag and a "
        "directory not visible through a normal listing.",
        "Clues within that hidden directory point toward an old backup location elsewhere "
        "on the system.",
        "Investigating the backup location reveals a third encoded flag and a private key "
        "belonging to a second, more privileged account.",
        "That key grants SSH access to the second account, which is deliberately more "
        "privileged than the first but is still not fully privileged.",
        "Reviewing that account's command history reveals a fourth encoded flag.",
        "Further investigation reveals a narrowly-scoped, intentionally controlled path that "
        "permits execution of one specific administrative utility - a custom binary - under "
        "tightly constrained conditions.",
        "Participants transfer that utility off the target system and analyze it using "
        "reverse-engineering tooling (such as Ghidra) to recover a required value.",
        "Supplying that recovered value through the permitted, controlled path results in "
        "full administrative access to the target system.",
        "The fifth and final flag, also stored in encoded form, is located within the "
        "administrator's own area of the system once full access is achieved.",
    ])

    # ---------------- Five-Flag Progression ----------------
    heading(doc, "7. Five-Flag Progression")
    para(doc,
         "Each flag marks the successful completion of a distinct stage, with purpose, "
         "objective, and tested skill deliberately escalating as the exercise progresses.")
    add_figure(doc, "flag_progression.png", "Figure 2: The five stages of the exercise")
    add_table(doc, ["Flag", "Stage", "Purpose", "Skill Tested"], [
        ["Flag 1", "Web Reconnaissance",
         "Establishes the reconnaissance mindset and rewards thorough web enumeration.",
         "Web recon, crawler-control file analysis, encoding awareness"],
        ["Flag 2", "Initial Access",
         "Rewards successful credential recovery and first authenticated foothold.",
         "Steganography, wordlist-based recovery, SSH usage"],
        ["Flag 3", "Lateral Discovery",
         "Encourages thorough host-level enumeration beyond the obvious.",
         "Filesystem enumeration, permissions analysis, artifact discovery"],
        ["Flag 4", "Secondary Access",
         "Rewards pivoting to a second, more privileged account.",
         "Key-based SSH access, activity-history review"],
        ["Flag 5", "Full Compromise",
         "Culminates the exercise via a deliberately controlled escalation path.",
         "Reverse engineering, controlled privilege-escalation reasoning"],
    ], col_widths=[0.7, 1.5, 2.6, 2.4])

    # ---------------- Rabbit Holes ----------------
    heading(doc, "8. Rabbit Holes")
    para(doc,
         "Three intentional rabbit holes are built into the environment. Each is designed "
         "to be believable enough that a participant may reasonably spend time "
         "investigating it, which reinforces the value of careful verification over "
         "assumption. None of the three provides an actual flag, grants unintended "
         "privileged access, or offers a shortcut past the intended path - each is a "
         "deliberate, safe dead end.")
    add_table(doc, ["Category", "Purpose"], [
        ["Decoy data location", "Tests whether participants verify content rather than assuming any "
         "plausible-looking location is the correct one."],
        ["Decoy web element", "Tests discipline against investigating superficially interesting but "
         "ultimately unproductive web content."],
        ["Decoy system artifact", "Tests whether participants recognize the difference between something "
         "that looks exploitable and something that actually is."],
    ], col_widths=[2.2, 5.0])
    note_box(doc, "Exact rabbit hole locations and content are documented separately for the "
                  "engagement facilitator only, and are not included in this proposal.")

    # ---------------- Hint/Riddle System ----------------
    heading(doc, "9. Hint/Riddle System")
    para(doc,
         "Every major stage has an optional hint available on request from the facilitator. "
         "Hints are delivered as riddles rather than direct instructions, so that a "
         "participant who requests one still has to do the work of interpreting and "
         "applying it, rather than simply being handed the next step.")
    note_box(doc, "Illustrative example only (not an actual hint used in the exercise): "
                  "“A door does not always keep its key hanging beside it.”")

    # ---------------- Skills Tested ----------------
    heading(doc, "10. Skills Tested")
    add_table(doc, ["Category", "Representative Skills"], [
        ["Reconnaissance", "Network/service discovery, web enumeration"],
        ["Web analysis", "Crawler-control file review, content inspection, encoding recognition"],
        ["Steganography", "Hidden-data detection and extraction, wordlist-based passphrase recovery"],
        ["Access & credentials", "SSH usage (password and key-based), credential decoding"],
        ["Linux enumeration", "Filesystem and permissions analysis, artifact discovery, log/history review"],
        ["Reverse engineering", "Static analysis of a compiled utility using standard tooling"],
        ["Privilege escalation", "Recognizing and correctly using a deliberately constrained escalation path"],
    ], col_widths=[2.2, 5.0])

    # ---------------- Difficulty Progression ----------------
    heading(doc, "11. Difficulty Progression")
    para(doc,
         "The exercise is calibrated to a medium-to-hard overall difficulty band, with each "
         "flag building on skills and context established by the one before it. Early "
         "stages are approachable to participants with foundational offensive-security "
         "knowledge, while later stages require sustained investigation and the ability to "
         "combine multiple techniques.")
    add_figure(doc, "difficulty_curve.png", "Figure 3: Relative difficulty by stage", width=5.4)

    # ---------------- Participant Experience ----------------
    heading(doc, "12. Participant Experience")
    para(doc,
         "Participants are given only the CTF target's IP address and a short briefing "
         "covering the storyline, objective, flag format, ground rules, and how to request "
         "a hint. Everything else - services, entry points, accounts, and the path forward "
         "- is discovered through the participant's own work.")
    bullets(doc, [
        "No installation or setup is required on the participant side beyond a standard "
        "penetration-testing distribution (Kali Linux or Parrot Security).",
        "The exercise can be run individually or in teams, and can be reset between "
        "sessions to give every participant or team an identical starting state.",
        "The environment is non-destructive: participant actions are contained entirely "
        "within the isolated exercise network.",
        "Facilitators can pause the event to issue riddle-based hints without revealing "
        "solutions outright.",
    ])

    # ---------------- Security/Isolation Considerations ----------------
    heading(doc, "13. Security/Isolation Considerations")
    bullets(doc, [
        "The entire exercise runs inside a dedicated virtual environment with no route to "
        "the client's production systems, management interfaces, or other virtual machines.",
        "No real client credentials, keys, or data are used anywhere in the exercise - every "
        "credential, key, and flag is generated specifically for the exercise environment.",
        "The target and participant networks are logically separated from the underlying "
        "virtualization host and from each other's broader environment.",
        "The environment is fully resettable, so it can be safely reused across multiple "
        "sessions without any risk of state or credentials carrying over.",
        "No component of the exercise requires ongoing internet access once deployed.",
    ])

    # ---------------- Proposed Deliverables ----------------
    heading(doc, "14. Proposed Deliverables")
    add_table(doc, ["Deliverable", "Description"], [
        ["Configured CTF environment", "A fully built and validated target system, ready to deploy within "
         "the client's Proxmox infrastructure."],
        ["Architecture & isolation guidance", "Documentation covering the recommended network layout and isolation "
         "approach for the exercise."],
        ["Facilitator operations guide", "A separate, organizer-only guide covering how to run, reset, and "
         "support the exercise during an event."],
        ["Participant briefing materials", "A short, spoiler-free briefing document for distribution to "
         "participants at the start of the event."],
        ["Reset & re-use capability", "The ability to reset the environment to a clean state between "
         "sessions, teams, or cohorts."],
    ], col_widths=[2.3, 4.8])

    # ---------------- Supplementary Practice Environments ----------------
    heading(doc, "15. Supplementary Practice Environments (Optional)")
    para(doc,
         "In addition to the core five-flag exercise, we can optionally provide a set of "
         "well-established, industry-standard vulnerable-by-design applications as separate, "
         "independent practice VMs. These are not part of the main storyline or attack path - "
         "they run as standalone, isolated machines alongside the core CTF target - and give "
         "participants a lower-pressure space to build and reinforce foundational skills "
         "before or in parallel with the main exercise.")
    add_table(doc, ["Practice VM", "Focus Area"], [
        ["OWASP Juice Shop", "Modern web application vulnerabilities across the OWASP Top 10, "
         "including API security and business-logic flaws in a realistic e-commerce application."],
        ["DVWA (Damn Vulnerable Web Application)", "Foundational, configurable-difficulty web "
         "application vulnerabilities - a strong on-ramp for participants newer to web security."],
        ["Metasploitable", "Intentionally vulnerable Linux server for service enumeration, exploitation "
         "practice, and hands-on familiarity with the Metasploit framework."],
    ], col_widths=[2.6, 4.5])
    para(doc,
         "Each practice VM is deployed and isolated the same way as the main CTF target, with "
         "no route between them or to production infrastructure. Because these are widely used, "
         "publicly available training platforms rather than custom-built environments, they add "
         "broad additional skill coverage - particularly classic web-application vulnerability "
         "classes - at minimal additional setup effort.")
    note_box(doc,
             "These practice VMs are entirely optional and are offered as a supplementary "
             "add-on. They do not affect the core five-flag storyline, its difficulty, or its "
             "flags in any way.")

    # ---------------- Conclusion ----------------
    heading(doc, "16. Conclusion")
    para(doc,
         "This exercise offers a realistic, engaging, and technically comprehensive way to "
         "exercise offensive-security skills in a controlled, safe, and fully isolated "
         "environment. Its progressive five-flag structure, believable storyline, and "
         "deliberately controlled privilege-escalation path make it suitable for skills "
         "development, team exercises, or assessment purposes, without introducing any risk "
         "to the client's production environment.")
    para(doc,
         "We welcome the opportunity to discuss scope, scheduling, and any customization "
         "the client would like to see reflected in the final exercise.", italic=True, color=MUTED)

    out_path = os.path.join(OUT_DIR, "CTF_Client_Proposal.docx")
    doc.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    build()
