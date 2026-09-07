"""
Converts the project's .md documentation into styled .docx files, embedding
the generated diagrams (see diagrams/gen_diagrams.py) at the right spots.

Usage: python3 build_docx.py
Requires: python-docx (pip install python-docx). No internet access needed.
"""
import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIAGRAMS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "diagrams")
OUT_DIR = os.path.dirname(os.path.abspath(__file__))

ACCENT = RGBColor(0x7A, 0x1F, 0x2B)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5B, 0x64, 0x72)

DIAGRAM_PLACEMENT = {
    "INSTALLATION.md": [
        ("4. Install", "install_workflow.png", "Figure: install_ctf.sh step-by-step workflow"),
    ],
    "ORGANIZER_GUIDE.md": [
        ("Intended attack path", "attack_path.png", "Figure: the intended player path, flag by flag"),
        ("Users", "user_privilege.png", "Figure: privilege boundaries between Agent99, Agent999, and root"),
    ],
    "ORGANIZER_SOLUTIONS.md": [
        ("Organizer Solutions", "attack_path.png", "Figure: the intended player path, flag by flag"),
    ],
    "NETWORK_GUIDE.md": [
        ("Network isolation", "network_topology.png", "Figure: network isolation boundaries"),
    ],
    "SECURITY.md": [
        ("Everything in this list is intentional", "user_privilege.png", "Figure: privilege boundaries between Agent99, Agent999, and root"),
    ],
    "RESET_GUIDE.md": [
        ("Reset Guide", "reset_workflow.png", "Figure: choosing a reset option"),
    ],
}

DOC_ORDER = [
    "README.md",
    "INSTALLATION.md",
    "ORGANIZER_GUIDE.md",
    "ORGANIZER_SOLUTIONS.md",
    "NETWORK_GUIDE.md",
    "SECURITY.md",
    "TROUBLESHOOTING.md",
    "RESET_GUIDE.md",
    "PLAYER_GUIDE.md",
]

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
STOP_MARKERS = re.compile(r"^(\d+\.\s|#{1,3}\s|-\s|\||>|---\s*$)")


def is_fence(line):
    return line.strip().startswith("```")


def clean_inline_markers(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


# ---------------------------------------------------------------------------
def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)

    title = doc.styles["Title"]
    title.font.name = "Calibri Light"
    title.font.size = Pt(30)
    title.font.color.rgb = ACCENT

    for lvl, size in ((1, 18), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.font.bold = True
        st.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(6)

    try:
        code_style = doc.styles.add_style("CodeBlock", 1)
    except ValueError:
        code_style = doc.styles["CodeBlock"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9.5)
    code_style.font.color.rgb = INK
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(10)
    code_style.paragraph_format.left_indent = Cm(0.4)


def shade_paragraph(paragraph, hex_color="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B8AFA0")
    borders.append(bottom)
    pPr.append(borders)
    p.paragraph_format.space_after = Pt(14)


CODE_SPAN_RE = re.compile(r"`([^`]+)`")


def emit_text_with_code_spans(paragraph, text, bold=False):
    """Emits text that may itself contain `code` spans (handles bold-wrapping-code)."""
    pos = 0
    for m in CODE_SPAN_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.bold = bold
        r = paragraph.add_run(m.group(1))
        r.bold = bold
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r.font.color.rgb = ACCENT
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = bold


def add_inline_runs(paragraph, text, doc_basename_map):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            emit_text_with_code_spans(paragraph, token[2:-2], bold=True)
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = ACCENT
        else:
            lm = LINK_RE.match(token)
            link_text, target = lm.group(1), lm.group(2)
            target_file = target.split("#")[0]
            if target_file in doc_basename_map:
                add_hyperlink(paragraph, link_text, doc_basename_map[target_file], external=True)
            elif target.startswith("http"):
                add_hyperlink(paragraph, link_text, target, external=True)
            else:
                r = paragraph.add_run(link_text)
                r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_hyperlink(paragraph, text, target, external=False):
    part = paragraph.part
    r_id = part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=external,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5C99")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def parse_table(lines):
    rows = [l.strip().strip("|").split("|") for l in lines if not re.match(r"^\|?\s*-{2,}", l.replace(" ", ""))]
    return [[c.strip() for c in row] for row in rows]


def emit_code_block(doc, code_lines, indent=False):
    p = doc.add_paragraph(style="CodeBlock")
    shade_paragraph(p)
    if indent:
        p.paragraph_format.left_indent = Cm(1.0)
    for idx, cl in enumerate(code_lines):
        if idx > 0:
            p.add_run().add_break()
        p.add_run(cl if cl else " ")


def consume_code_fence(lines, i):
    base_indent = len(lines[i]) - len(lines[i].lstrip())
    j = i + 1
    code_lines = []
    while j < len(lines) and not is_fence(lines[j]):
        l = lines[j]
        if l[:base_indent].strip() == "":
            l = l[base_indent:]
        code_lines.append(l.rstrip("\n"))
        j += 1
    j += 1  # skip closing fence
    return code_lines, j


def consume_paragraph_text(lines, i):
    text_parts = [lines[i].strip()]
    j = i + 1
    while j < len(lines):
        s = lines[j].strip()
        if s == "" or STOP_MARKERS.match(s) or is_fence(lines[j]):
            break
        text_parts.append(s)
        j += 1
    return " ".join(text_parts), j


def consume_list_item(lines, i, doc, doc_basename_map, style, strip_prefix_re):
    text, i = consume_paragraph_text(lines, i)
    text = strip_prefix_re.sub("", text, count=1)
    p = doc.add_paragraph(style=style)
    add_inline_runs(p, text, doc_basename_map)

    while i < len(lines):
        s = lines[i].strip()
        if is_fence(lines[i]):
            code_lines, i = consume_code_fence(lines, i)
            emit_code_block(doc, code_lines, indent=True)
        elif s and not STOP_MARKERS.match(s):
            more_text, i = consume_paragraph_text(lines, i)
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = Cm(0.6)
            add_inline_runs(p2, more_text, doc_basename_map)
        else:
            break
    return i


# ---------------------------------------------------------------------------
def build_doc(md_filename, doc_basename_map):
    src = os.path.join(PROJECT_ROOT, md_filename)
    with open(src, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    set_base_styles(doc)

    placements = DIAGRAM_PLACEMENT.get(md_filename, [])
    placed = set()

    i = 0
    while i < len(lines):
        line = lines[i]

        if is_fence(line):
            code_lines, i = consume_code_fence(lines, i)
            emit_code_block(doc, code_lines)
            continue

        if line.startswith("# "):
            text = clean_inline_markers(line[2:].strip())
            doc.add_heading(text, level=0)
            for after_text, img, caption in placements:
                key = (after_text, img)
                if key not in placed and after_text.lower() in text.lower():
                    insert_image_after(doc, img, caption)
                    placed.add(key)
            i += 1
            continue

        if line.startswith("## "):
            text = clean_inline_markers(line[3:].strip())
            doc.add_heading(text, level=1)
            for after_text, img, caption in placements:
                key = (after_text, img)
                if key not in placed and after_text.lower() in text.lower():
                    insert_image_after(doc, img, caption)
                    placed.add(key)
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(clean_inline_markers(line[4:].strip()), level=2)
            i += 1
            continue

        if line.strip() == "---":
            add_hr(doc)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table(table_lines)
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                try:
                    table.style = "Light Grid Accent 2"
                except KeyError:
                    table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.LEFT
                for r_idx, row in enumerate(rows):
                    cells = table.add_row().cells
                    for c_idx, cell_text in enumerate(row):
                        if c_idx >= len(cells):
                            continue
                        p = cells[c_idx].paragraphs[0]
                        add_inline_runs(p, cell_text, doc_basename_map)
                        if r_idx == 0:
                            for run in p.runs:
                                run.bold = True
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if re.match(r"^\d+\.\s", line):
            i = consume_list_item(lines, i, doc, doc_basename_map, "List Number", re.compile(r"^\d+\.\s+"))
            continue

        if line.startswith("- "):
            i = consume_list_item(lines, i, doc, doc_basename_map, "List Bullet", re.compile(r"^-\s+"))
            continue

        if line.strip() == "":
            i += 1
            continue

        text, i = consume_paragraph_text(lines, i)
        p = doc.add_paragraph()
        add_inline_runs(p, text, doc_basename_map)

    out_path = os.path.join(OUT_DIR, md_filename.replace(".md", ".docx"))
    doc.save(out_path)
    print("wrote", out_path)


def insert_image_after(doc, img_filename, caption):
    img_path = os.path.join(DIAGRAMS, img_filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    cap.paragraph_format.space_after = Pt(14)


def main():
    doc_basename_map = {name: name.replace(".md", ".docx") for name in DOC_ORDER}
    for name in DOC_ORDER:
        build_doc(name, doc_basename_map)


if __name__ == "__main__":
    main()
